import pandas as pd
import numpy as np
import openpyxl
import streamlit as st
import io

st.set_page_config(page_title="Battery Data Processor", layout="wide")
st.title("Universal Battery Data Processor")
st.write("Upload your raw test data to generate the processed Excel file and charts for DCA, DoD 17.5%, and DoD 50% tests.")

# --- Step 1: Input Setup ---
uploaded_file = st.file_uploader("Select Raw Data Excel File", type=["xlsx"])

if uploaded_file is not None:
    with st.spinner("Analyzing and Processing Data..."):
        try:
            # --- Step 2: Test Type & Sheet Detection ---
            xl = pd.ExcelFile(uploaded_file, engine='openpyxl')
            all_sheets = xl.sheet_names

            # --- Step 3: Filtering of Sheet ---
            target_sheet = None
            test_type = None
            is_failed = False

            for s in all_sheets:
                s_lower = s.lower()
                
                if "failed" in s_lower:
                    is_failed = True
                    
                if "17.5" in s_lower:
                    test_type = "DoD_17.5"
                    target_sheet = s
                    break

                elif "50" in s_lower:
                    test_type = "DoD_50"
                    target_sheet = s
                    break

                elif "dca" in s_lower and 'flooded' not in s_lower:
                    test_type = "DCA"
                    target_sheet = s
                    break
                    
                elif "mht" in s_lower:
                    test_type = "MHT"
                    target_sheet = s
                    break

            if not target_sheet:
                st.error("No valid test sheet was identified in the workbook.")
                st.stop()

            # --- Step 4: Read Raw Data ---
            df_meta = pd.read_excel(xl, sheet_name=target_sheet, header=None, nrows=4)
            if df_meta.shape[0] >= 4 and df_meta.shape[1] >= 2:
                battery_code = str(df_meta.iloc[3, 1]).strip()
            else:
                battery_code = "UNKNOWN"

            raw_df = pd.read_excel(xl, sheet_name=target_sheet, skiprows=29)
            raw_df.columns = raw_df.columns.astype(str).str.strip()

            # --- CRITICAL FIX: Standardize Step Time Format ---
            def standardize_time(val):
                if pd.isna(val):
                    return pd.NaT
                val_str = str(val).strip()
                # If hours are missing (e.g., '00:00.000'), append '00:' to match hh:mm:ss.ttt
                if val_str.count(':') == 1:
                    val_str = '00:' + val_str
                return pd.to_timedelta(val_str, errors='coerce')
                
            if 'Step Time' in raw_df.columns:
                raw_df['Step Time'] = raw_df['Step Time'].apply(standardize_time)

            status_text = " - Failed" if is_failed else ""
            output_filename = f'Processed {test_type}{status_text} Battery Results of {battery_code}.xlsx'

            # --- Create the Output File in Memory ---
            output_buffer = io.BytesIO()

            # --- Step 5: Execute Test ---
            with pd.ExcelWriter(output_buffer, engine='xlsxwriter') as writer:
                workbook = writer.book
                
                # Test A: DCA TEST ENGINE
                if test_type == "DCA":
                    # --- IC Sheet Filtering and Computation ---
                    df = raw_df.copy()
                    df = df[df['Status'] == 'CHA']
                    df = df[df['Step'].isin([30])]
                    df = df[df['Voltage'] >= 14.8]
                    df = df[df['Current'] >= 0]
                    df = df[(df['Step Time'] >= pd.to_timedelta('00:00:10.000')) & (df['Step Time'] <= pd.to_timedelta('00:00:10.100'))]
                    df = df.dropna(subset=['Temperature'])
                    df['Ampere'] = df['AhStep'] * 360
                    df['A/Ah'] = df['Ampere'] / 60

                    total_ah1 = df['AhStep'].sum()
                    total_as1 = total_ah1 * 3600
                    time_sec1 = 200
                    amp1 = total_as1 / time_sec1
                    capacity1 = 60
                    a_ah1 = amp1 / capacity1
                    factor1 = 0.512
                    net1 = a_ah1 * factor1

                    df_summary = pd.DataFrame({
                        'Metric': ['Total-Ah', 'Total-AS', 'Time-Sec', 'Amp', 'Capacity-Ah', 'A/Ah', 'Factor', 'Net'],
                        'Value':  [total_ah1, total_as1, time_sec1, amp1, capacity1, a_ah1, factor1, net1]
                    })

                    # --- ID Sheet Filtering and Computation ---
                    df2 = raw_df.copy()
                    df2 = df2[df2['Status'] == 'CHA']
                    df2 = df2[df2['Step'].isin([43])]
                    df2 = df2[df2['Voltage'] >= 14.8]
                    df2 = df2[df2['Current'] >= 0]
                    df2 = df2[(df2['Step Time'] >= pd.to_timedelta('00:00:10.000')) & (df2['Step Time'] <= pd.to_timedelta('00:00:10.100'))]
                    df2 = df2.dropna(subset=['Temperature'])
                    df2['Ampere'] = df2['AhStep'] * 360
                    df2['A/Ah'] = df2['Ampere'] / 60

                    total_ah2 = df2['AhStep'].sum()
                    total_as2 = total_ah2 * 3600
                    time_sec2 = 200
                    amp2 = total_as2 / time_sec2
                    capacity2 = 60
                    a_ah2 = amp2 / capacity2
                    factor2 = 0.223
                    net2 = a_ah2 * factor2

                    df2_summary = pd.DataFrame({
                        'Metric': ['Total-Ah', 'Total-AS', 'Time-Sec', 'Amp', 'Capacity-Ah', 'A/Ah', 'Factor', 'Net'],
                        'Value':  [total_ah2, total_as2, time_sec2, amp2, capacity2, a_ah2, factor2, net2]
                    })

                    # --- IR Sheet and Computation ---
                    df3 = raw_df.copy()
                    df3 = df3[df3['Status'] == 'CHA']
                    df3 = df3[df3['Step'].isin([70, 81])]
                    df3 = df3[df3['Cycle Level'] == 3]
                    df3 = df3[(df3['Step Time'] >= pd.to_timedelta('00:00:05.000')) & (df3['Step Time'] <= pd.to_timedelta('00:00:05.100'))]
                    df3 = df3.dropna(subset=['Temperature'])
                    
                    df3['Time Stamp'] = df3['Time Stamp'].astype(str).str.strip()
                    df3['Status'] = df3['Status'].astype(str).str.strip()
                    df3['Step'] = df3['Step'].astype(str).str.strip()
                    df3 = df3.drop_duplicates(subset=['Time Stamp', 'Step', 'Status'], keep='first')
                    
                    df3['Ampere'] = df3['AhStep'] * 720
                    df3['A/Ah'] = df3['Ampere'] / 60

                    # --- Tracker --- 
                    tracker = (df3['Cycle'] < df3['Cycle'].shift(1)).cumsum() + 1
                    df3['S'] = 'S' + tracker.astype(str)

                    # --- IR Calculation ---
                    IR_Calculation = df3.groupby('S', sort=False)['AhStep'].sum().reset_index()
                    IR_Calculation = IR_Calculation.rename(columns={'AhStep': 'Total AhStep'})

                    grand_total_ah = df3['AhStep'].sum()
                    total_as3 = grand_total_ah * 3600
                    time_sec3 = 2850
                    amp3 = total_as3 / time_sec3
                    capacity3 = 60
                    a_ah3 = amp3 / capacity3
                    factor3 = 0.218
                    net3 = a_ah3 * factor3

                    df4 = pd.DataFrame({
                        '1': ['Total-Ah', 'Total-AS', 'Time-Sec', 'Amp', 'Capacity-Ah', 'A/Ah', 'Factor', 'Net'],
                        '2': [grand_total_ah, total_as3, time_sec3, amp3, capacity3, a_ah3, factor3, net3]
                    })

                    Ic, Id, Ir, Constant = net1, net2, net3, 0.181
                    final_a_ah = Ic + Id + Ir - Constant

                    df5 = pd.DataFrame({
                        '1': ['Ic', 'Id', 'Ir', 'Constant', 'Final A/Ah'],
                        '2': [Ic, Id, Ir, Constant, final_a_ah]
                    })

                    # --- Normalized, Standardized, OCV ---
                    target_rows = 55
                    df4_filtered = df3.groupby(tracker)['A/Ah'].mean()
                    dca_combined = pd.concat([df['A/Ah'], df2['A/Ah'], df4_filtered]).tolist()

                    OCV_filtered_81 = raw_df[(raw_df['Step'].astype(str).str.strip() == '81') & (raw_df['Current'] == 0) & (raw_df['Temperature'].notna()) & (raw_df['Cycle'].astype(str).str.strip() == '19')]['Voltage'].dropna()
                    OCV_filtered_30 = raw_df[(raw_df['Step'].astype(str).str.strip() == '30') & (raw_df['Current'] == 0) & (raw_df['Temperature'].notna())]['Voltage'].dropna()
                    OCV_filtered_43 = raw_df[(raw_df['Step'].astype(str).str.strip() == '43') & (raw_df['Current'] == 0) & (raw_df['Temperature'].notna())]['Voltage'].dropna()

                    OCV_Combined = pd.concat([OCV_filtered_30, OCV_filtered_43, OCV_filtered_81]).reset_index(drop=True).iloc[:target_rows].tolist()
                    normalized = [final_a_ah] * target_rows

                    df6 = pd.DataFrame({
                        'Normalized': normalized,
                        'Standard DCA': dca_combined[:target_rows],
                        'OCV': OCV_Combined
                    })

                    # --- Excel Output ---
                    df.to_excel(writer, sheet_name='IC', index=False)
                    df2.to_excel(writer, sheet_name='ID', index=False)
                    df3.to_excel(writer, sheet_name='IR', index=False)
                    IR_Calculation.to_excel(writer, sheet_name='Final', index=False)
                    df4.to_excel(writer, sheet_name='Final', index=False, header=False, startcol=len(IR_Calculation.columns) + 1)
                    df5.to_excel(writer, sheet_name='Final', index=False, header=False, startcol=len(IR_Calculation.columns) + 4)
                    df6.to_excel(writer, sheet_name='Cumulative', index=False)

                    worksheet1 = writer.sheets['IC']
                    worksheet2 = writer.sheets['ID']
                    worksheet3 = writer.sheets['IR']
                    worksheet_final = writer.sheets['Final']
                    worksheet_cum = writer.sheets['Cumulative']

                    # --- Summary of each Sheet --- 
                    yellow_format = workbook.add_format({'bg_color': '#FFFF00'}) #[cite: 1]
                    
                    df3_reset = df3.reset_index(drop=True) #[cite: 1]
                    first_row_indices = df3_reset.groupby('S').head(1).index.tolist() #[cite: 1]
                    last_col = len(df3.columns) - 1 #[cite: 1]
                    
                    for idx in first_row_indices: #[cite: 1]
                        excel_row = idx + 1   #[cite: 1]
                        worksheet3.conditional_format(excel_row, 0, excel_row, last_col, { #[cite: 1]
                            'type': 'formula', #[cite: 1]
                            'criteria': '=TRUE', #[cite: 1]
                            'format': yellow_format #[cite: 1]
                        }) #[cite: 1]

                    worksheet_final.conditional_format('A1:B16', {'type': 'no_blanks', 'format': workbook.add_format({'bg_color':'#FFF2CC'})})
                    worksheet_final.conditional_format('D1:E8',  {'type': 'no_blanks', 'format': workbook.add_format({'bg_color': '#E2EFDA'})})
                    worksheet_final.conditional_format('G1:H5',  {'type': 'no_blanks', 'format': workbook.add_format({'bg_color': '#D9E1F2'})})

                    orange_format = workbook.add_format({'bg_color': '#FFC000', 'font_color': '#000000'})
                    orange_num_format = workbook.add_format({'num_format': '0.0000', 'bg_color': '#FFC000', 'font_color': '#000000'})

                    for i, row in df_summary.iterrows():
                        worksheet1.write(1 + i, 21, row['Metric'], orange_format)
                        worksheet1.write(1 + i, 22, row['Value'], orange_num_format)

                    for i, row in df2_summary.iterrows():
                        worksheet2.write(1 + i, 21, row['Metric'], orange_format)
                        worksheet2.write(1 + i, 22, row['Value'], orange_num_format)

                    # --- DCA Charting Layout ---
                    chart1 = workbook.add_chart({'type': 'line'})
                    chart1.add_series({
                        'categories': ['IC', 1, df.columns.get_loc('Cycle'), len(df), df.columns.get_loc('Cycle')],
                        'values':     ['IC', 1, df.columns.get_loc('Current'), len(df), df.columns.get_loc('Current')],
                        'line': {'width': 1.0, 'color': 'blue'},
                        'marker': {'type': 'circle', 'size': 5, 'border': {'color': 'blue'}, 'fill': {'color': 'blue'}}})
                    c1_diff = df['Current'].max() - df['Current'].min()
                    chart1.set_title({'name': f'{battery_code} - Ic (After Charge History)'}) #[cite: 1]
                    chart1.set_x_axis({'name': '20 Micro cycles', 'label_position': 'low'})
                    chart1.set_y_axis({'name': 'Charge Current (A)', 'min': df['Current'].min() - (c1_diff/5), 'max': df['Current'].max() + (c1_diff/5)})
                    chart1.set_size({'width': 750, 'height': 300})
                    worksheet1.insert_chart('B24', chart1)

                    chart2 = workbook.add_chart({'type': 'line'})
                    chart2.add_series({
                        'categories': ['IC', 1, df.columns.get_loc('Cycle'), len(df), df.columns.get_loc('Cycle')],
                        'values':     ['IC', 1, df.columns.get_loc('A/Ah'), len(df), df.columns.get_loc('A/Ah')],
                        'line': {'width': 1.0, 'color': 'blue'},
                        'marker': {'type': 'circle', 'size': 5, 'border': {'color': 'blue'}, 'fill': {'color': 'blue'}}})
                    c2_diff = df['A/Ah'].max() - df['A/Ah'].min()
                    chart2.set_title({'name': f'{battery_code} - Ic Normalized (After Charge History)'}) #[cite: 1]
                    chart2.set_x_axis({'name': '20 Micro cycles', 'label_position': 'low'}) #[cite: 1]
                    chart2.set_y_axis({'name': 'Normalized Charge Current (A/Ah)', 'min': df['A/Ah'].min() - (c2_diff/5), 'max': df['A/Ah'].max() + (c2_diff/5)})
                    chart2.set_size({'width': 750, 'height': 300})
                    worksheet1.insert_chart('N24', chart2)

                    chart3 = workbook.add_chart({'type': 'line'})  
                    chart3.add_series({
                        'categories': ['ID', 1, df2.columns.get_loc('Cycle'), len(df2), df2.columns.get_loc('Cycle')],
                        'values':     ['ID', 1, df2.columns.get_loc('Current'), len(df2), df2.columns.get_loc('Current')],
                        'line': {'width': 1.0, 'color': 'blue'},
                        'marker': {'type': 'circle', 'size': 5, 'border': {'color': 'blue'}, 'fill': {'color': 'blue'}}})
                    c3_diff = df2['Current'].max() - df2['Current'].min()
                    chart3.set_title({'name': f'{battery_code} - Id (After Discharge History)'})
                    chart3.set_x_axis({'name': '20 Micro cycles', 'label_position': 'low'})
                    chart3.set_y_axis({'name': 'Charge Current A', 'min': df2['Current'].min() - (c3_diff/5), 'max': df2['Current'].max() + (c3_diff/5)})
                    chart3.set_size({'width': 750, 'height': 300})
                    worksheet2.insert_chart('B24', chart3)

                    chart4 = workbook.add_chart({'type': 'line'}) 
                    chart4.add_series({
                        'categories': ['ID', 1, df2.columns.get_loc('Cycle'), len(df2), df2.columns.get_loc('Cycle')],
                        'values':     ['ID', 1, df2.columns.get_loc('A/Ah'), len(df2), df2.columns.get_loc('A/Ah')],
                        'line': {'width': 1.0, 'color': 'blue'},
                        'marker': {'type': 'circle', 'size': 5, 'border': {'color': 'blue'}, 'fill': {'color': 'blue'}}})
                    c4_diff = df2['A/Ah'].max() - df2['A/Ah'].min()
                    chart4.set_title({'name': f'{battery_code} - Id Normalized (After Discharge History)'})
                    chart4.set_x_axis({'name': '20 Micro cycles', 'label_position': 'low'})
                    chart4.set_y_axis({'name': 'Normalized Charge Current (A/Ah)', 'min': df2['A/Ah'].min() - (c4_diff/5), 'max': df2['A/Ah'].max() + (c4_diff/5)})
                    chart4.set_size({'width': 750, 'height': 300})
                    worksheet2.insert_chart('N24', chart4)

                    chart_ir = workbook.add_chart({'type': 'scatter', 'subtype': 'smooth_with_markers'})
                    cycle_length = 38 
                    total_ir_rows = len(df3)
                    ir_cycle_col = df3.columns.get_loc('Cycle')
                    ir_current_col = df3.columns.get_loc('Current')

                    cycle_num = 1
                    for s_row in range(1, total_ir_rows + 1, cycle_length):
                        e_row = min(s_row + cycle_length - 1, total_ir_rows)
                        chart_ir.add_series({
                            'name':       f'{cycle_num}', #[cite: 1]
                            'categories': ['IR', s_row, ir_cycle_col, e_row, ir_cycle_col],
                            'values':     ['IR', s_row, ir_current_col, e_row, ir_current_col],
                            'marker':     {'type': 'circle', 'size': 4},
                            'line':       {'width': 1.0}
                        })
                        cycle_num += 1
                    chart_ir.set_title({'name': f'{battery_code} - IR Real World Start-Stop Regenerative Current'})
                    chart_ir.set_x_axis({'name': 'Cycle', 'min': 0, 'max': 20, 'major_unit': 1})
                    chart_ir.set_y_axis({'name': 'Charge Current (A)'}) #[cite: 1]
                    chart_ir.set_size({'width': 1000, 'height': 500})
                    worksheet3.insert_chart('V2', chart_ir)

                    # --- Added Normalized DCA Chart ---
                    chart_ir_norm = workbook.add_chart({'type': 'scatter', 'subtype': 'smooth_with_markers'}) #[cite: 1]
                    ir_AAh_col = df3.columns.get_loc('A/Ah') #[cite: 1]

                    cycle_num = 1 #[cite: 1]
                    for s_row in range(1, total_ir_rows + 1, cycle_length): #[cite: 1]
                        e_row = min(s_row + cycle_length - 1, total_ir_rows) #[cite: 1]
                        chart_ir_norm.add_series({ #[cite: 1]
                            'name':       f'{cycle_num}', #[cite: 1]
                            'categories': ['IR', s_row, ir_cycle_col, e_row, ir_cycle_col], #[cite: 1]
                            'values':     ['IR', s_row, ir_AAh_col, e_row, ir_AAh_col], #[cite: 1]
                            'marker':     {'type': 'circle', 'size': 4}, #[cite: 1]
                            'line':       {'width': 1.0} #[cite: 1]
                        }) #[cite: 1]
                        cycle_num += 1 #[cite: 1]
                    chart_ir_norm.set_title({'name': f'{battery_code} - IR Real World Start-Stop Regenerative Current'}) #[cite: 1]
                    chart_ir_norm.set_x_axis({'name': 'Cycle', 'min': 0, 'max': 20, 'major_unit': 1}) #[cite: 1]
                    chart_ir_norm.set_y_axis({'name': 'Normalized DCA (A/Ah)', 'min': 0, 'max': 1.2, 'major_unit': 0.2}) #[cite: 1]
                    chart_ir_norm.set_size({'width': 1000, 'height': 500}) #[cite: 1]
                    worksheet3.insert_chart('V28', chart_ir_norm) #[cite: 1]

                    max_row = len(df6)
                    chart_cum = workbook.add_chart({'type': 'line'})
                    chart_cum.add_series({
                        'name':       ['Cumulative', 0, 0],
                        'values':     ['Cumulative', 1, 0, max_row, 0],
                        'line':       {'color': '#4472C4', 'width': 1.5},
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#4472C4'}, 'fill': {'color': '#4472C4'}}
                    })
                    chart_cum.add_series({
                        'name':       ['Cumulative', 0, 1],
                        'values':     ['Cumulative', 1, 1, max_row, 1],
                        'line':       {'none': True},
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#ED7D31'}, 'fill': {'color': '#ED7D31'}}
                    })
                    chart_cum.add_series({
                        'name':       ['Cumulative', 0, 2],
                        'values':     ['Cumulative', 1, 2, max_row, 2],
                        'line':       {'none': True},
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#70AD47'}, 'fill': {'color': '#70AD47'}},
                        'y2_axis':    True
                    })
                    chart_cum.set_title({'name': f'{battery_code} - Gen 3 DCA'}) 
                    chart_cum.set_y_axis({'name': 'Normalized Charge current (A/Ah)', 'min': 0.00, 'max': 1.10, 'major_unit': 0.10})
                    chart_cum.set_y2_axis({'name': 'OCV', 'min': 12.50, 'max': 13.30, 'major_unit': 0.10})
                    chart_cum.set_size({'width': 1000, 'height': 500})
                    chart_cum.set_legend({'none': True})
                    worksheet_cum.insert_chart('E2', chart_cum) #[cite: 1]

                # Test B: DoD 17.5% TEST ENGINE
                elif test_type == "DoD_17.5":
                    battery_name = battery_code
                    df = raw_df.copy()
                    df.rename(columns={df.columns[1]: 'Status'}, inplace=True)

                    # --- Calculate Voltage per Cell and Current per Cell ---
                    df['Voltage'] = pd.to_numeric(df['Voltage'], errors='coerce')
                    df['Current'] = pd.to_numeric(df['Current'], errors='coerce')
                    df['AhStep'] = pd.to_numeric(df['AhStep'], errors='coerce')

                    df['Voltage per Cell'] = df['Voltage'] / 6
                    df['Current per Cell'] = df['Current'] / 6
                    df['SoC'] = 0.00 

                    # --- SoC Computation ---
                    for i in range(len(df)):
                        if i == 0: 
                            first_ah = 0 if pd.isna(df.loc[i, 'AhStep']) else df.loc[i, 'AhStep']
                            if df.loc[i, 'Status'] == 'CHA':
                                df.loc[i, 'SoC'] = (60 + first_ah) / 60
                            else:
                                df.loc[i, 'SoC'] = (60 - first_ah) / 60
                        else: 
                            previous_SoC = df.loc[i-1, 'SoC']
                            ahstep = df.loc[i, 'AhStep']
                            ahstep_before = df.loc[i-1, 'AhStep']
                            diff_ahstep = ahstep - ahstep_before
                            
                            if pd.isna(ahstep) or pd.isna(ahstep_before) or ahstep == ahstep_before:
                                diff_ahstep = 0
                                
                            if df.loc[i, 'Status'] == 'CHA':
                                df.loc[i, 'SoC'] = previous_SoC + (diff_ahstep / 60)    
                            else:
                                df.loc[i, 'SoC'] = previous_SoC - (diff_ahstep / 60)
                                
                    df['SoC%'] = df['SoC'] / 1.0 * 100 
                    start_discharge = (df['SoC%'] < 100).idxmax() 
                    full_charge = (df.loc[start_discharge:, 'SoC%'] >= 100).idxmax() 
                    
                    if full_charge > start_discharge:
                        df = df.loc[:full_charge]

                    # --- Filtering Sheet 2 ---
                    df2 = raw_df.copy()
                    df2.rename(columns={df2.columns[1]: 'Status'}, inplace=True) 
                    df2 = df2[df2['Status'] == 'DCH']
                    df2 = df2[df2['Step'].isin([3, 7])]
                    # Update comparison to use pd.to_timedelta
                    df2 = df2[(df2['Step Time'] >= pd.to_timedelta('00:30:00.000')) & (df2['Step Time'] <= pd.to_timedelta('00:30:00.100'))]
                    
                    # --- Different Names for Temperatures --- 
                    temp_col = 'Temperature_1' if 'Temperature_1' in df2.columns else 'Temperature'
                    df2 = df2.dropna(subset=[temp_col])
                    
                    df2 = df2[df2['Cycle'] != df2['Cycle'].shift()]
                    df2['Total Cycle'] = range(0, len(df2)) 

                    # --- Excel Output ---
                    df.to_excel(writer, sheet_name='Sheet 1', index=False)
                    df2.to_excel(writer, sheet_name='Sheet 2', index=False)

                    worksheet1 = writer.sheets['Sheet 1']
                    worksheet2 = writer.sheets['Sheet 2']   

                    # --- Table for Total Cycle / 85 ---
                    max_total_cycle_calc = df2['Total Cycle'].max() / 85 #[cite: 1]
                    header_format = workbook.add_format({'bold': True, 'bg_color': '#D9E1F2', 'border': 1, 'align': 'center'}) #[cite: 1]
                    value_format = workbook.add_format({'border': 1, 'align': 'center', 'num_format': '0.00'}) #[cite: 1]
                    
                    worksheet2.write('U2', 'Total Cycle', header_format) #[cite: 1]
                    worksheet2.write('V2', max_total_cycle_calc, value_format) #[cite: 1]

                    # --- DoD17.5% Charting Layout ---
                    chart1 = workbook.add_chart({'type': 'line'})
                    chart1.add_series({
                        'categories': ['Sheet 1', 1, df.columns.get_loc('Prog Time'), len(df), df.columns.get_loc('Prog Time')],
                        'values':     ['Sheet 1', 1, df.columns.get_loc('Voltage per Cell'), len(df), df.columns.get_loc('Voltage per Cell')],
                        'line': {'width': 0.01, 'color': 'blue'}})
                    chart1.set_title({'name': f'{battery_name} - Voltage vs Time'})
                    chart1.set_x_axis({'name': 'Prog Time', 'label_position': 'low'})
                    chart1.set_y_axis({'name': 'Voltage per Cell', 'min': 1.5, 'max': 3})
                    chart1.set_size({'width': 1000, 'height': 500})
                    worksheet1.insert_chart('X2', chart1) #[cite: 1]

                    chart2 = workbook.add_chart({'type': 'line'})
                    chart2.add_series({
                        'categories': ['Sheet 1', 1, df.columns.get_loc('Prog Time'), len(df), df.columns.get_loc('Prog Time')],
                        'values':     ['Sheet 1', 1, df.columns.get_loc('Current per Cell'), len(df), df.columns.get_loc('Current per Cell')],
                        'line': {'width': 0.01, 'color': 'blue'}})
                    chart2.set_title({'name': f'{battery_name} - Current vs Time'})
                    chart2.set_x_axis({'name': 'Prog Time', 'label_position': 'low'})
                    chart2.set_y_axis({'name': 'Current per Cell', 'min': -4, 'max': 4})
                    chart2.set_size({'width': 1000, 'height': 500})
                    worksheet1.insert_chart('X28', chart2) #[cite: 1]

                    chart3 = workbook.add_chart({'type': 'line'})  
                    chart3.add_series({
                        'categories': ['Sheet 1', 1, df.columns.get_loc('Prog Time'), len(df), df.columns.get_loc('Prog Time')],
                        'values':     ['Sheet 1', 1, df.columns.get_loc('SoC%'), len(df), df.columns.get_loc('SoC%')],
                        'line': {'width': 0.01, 'color': 'blue'}})
                    chart3.set_title({'name': f'{battery_name} - State of Charge (SoC%) vs Time'})
                    chart3.set_x_axis({'name': 'Prog Time', 'label_position': 'low'})
                    chart3.set_y_axis({'name': 'SoC%', 'min': 0, 'max': 100})
                    chart3.set_size({'width': 1000, 'height': 500})
                    worksheet1.insert_chart('X54', chart3) #[cite: 1]

                    chart4 = workbook.add_chart({'type': 'scatter', 'subtype': 'straight_with_markers'})
                    chart4.add_series({
                        'categories': ['Sheet 2', 1, df2.columns.get_loc('Total Cycle'), len(df2), df2.columns.get_loc('Total Cycle')],
                        'values':     ['Sheet 2', 1, df2.columns.get_loc('Voltage'), len(df2), df2.columns.get_loc('Voltage')],
                        'line': {'width': 0.01, 'color': 'blue'}})
                    chart4.set_title({'name': f'17.5% DoD of {battery_name}'})  
                    chart4.set_x_axis({'name': 'Cycle', 'min': 0, 'label_position': 'low', 'major_unit': 85})
                    chart4.set_y_axis({'name': 'Voltage', 'min': 10, 'max': 12.5})
                    chart4.set_size({'width': 1000, 'height': 500})
                    worksheet2.insert_chart('U4', chart4) #[cite: 1]

                # Test C: DoD 50% TEST ENGINE
                elif test_type == "DoD_50":
                    # --- Dataframe 1 to get Voltage --- 
                    df = raw_df.copy() #[cite: 1]
                    df['Status'] = df['Status'].astype(str).str.strip() #[cite: 1]
                    df = df[df['Status'] == 'DCH'] #[cite: 1]
                    df['Step'] = pd.to_numeric(df['Step'], errors='coerce') #[cite: 1]
                    df = df[df['Step'] == 5] #[cite: 1]
                    df = df[(df['Step Time'] >= pd.to_timedelta('02:00:00.000')) & (df['Step Time'] <= pd.to_timedelta('02:00:00.100'))] #[cite: 1]

                    df = df.dropna(subset=['Temperature']) #[cite: 1]
                    df['Cycle'] = pd.to_numeric(df['Cycle'], errors='coerce') #[cite: 1]
                    df['Voltage'] = pd.to_numeric(df['Voltage'], errors='coerce') #[cite: 1]
                    df = df.drop_duplicates(subset=['Cycle'], keep='first') #[cite: 1]

                    # --- Dataframe 2 to get Charged Time --- 
                    df2 = raw_df.copy() #[cite: 1]
                    df2['Status'] = df2['Status'].astype(str).str.strip() #[cite: 1]
                    df2 = df2[df2['Status'].isin(['CHA', 'DCH'])] #[cite: 1]
                    df2['Step'] = pd.to_numeric(df2['Step'], errors='coerce') #[cite: 1]
                    df2 = df2[df2['Step'].isin([5,8])] #[cite: 1]

                    t_start2 = pd.to_timedelta('02:00:00.000') #[cite: 1]
                    t_end2 = pd.to_timedelta('02:05:00.000') #[cite: 1]
                    t_start3 = pd.to_timedelta('05:00:00.000') #[cite: 1]
                    t_end3 = pd.to_timedelta('05:00:00.100') #[cite: 1]

                    cond2 = (df2['Step Time'] >= t_start2) & (df2['Step Time'] <= t_end2) #[cite: 1]
                    cond3 = (df2['Step Time'] >= t_start3) & (df2['Step Time'] <= t_end3) #[cite: 1]

                    df2 = df2[cond2 | cond3] #[cite: 1]

                    first_cha = df2[df2['Status'] == 'CHA'].drop_duplicates(subset=['Cycle'], keep='first') #[cite: 1]
                    first_cha = first_cha[['Cycle', 'Prog Time']].rename(columns={'Prog Time': 'First_CHA_Time'}) #[cite: 1]

                    first_dch = df2[df2['Status'] == 'DCH'].drop_duplicates(subset=['Cycle'], keep='first') #[cite: 1]
                    first_dch = first_dch[['Cycle', 'Prog Time']].rename(columns={'Prog Time': 'First_DCH_Time'}) #[cite: 1]

                    calc_df = pd.merge(first_cha, first_dch, on='Cycle', how='outer').sort_values('Cycle') #[cite: 1]
                    calc_df['Next_Cycle_First_DCH'] = calc_df['First_DCH_Time'].shift(-1) #[cite: 1]

                    # --- Equation for Charged Time --- 
                    calc_df['Time in hrs'] = (
                        (pd.to_timedelta(calc_df['Next_Cycle_First_DCH']) - pd.to_timedelta(calc_df['First_CHA_Time']))
                        .dt.total_seconds() / 3600
                    ).round(2) #[cite: 1]

                    final_table = calc_df[['Cycle', 'Time in hrs']].dropna() #[cite: 1]

                    df2 = df2.dropna(subset=['Temperature']) #[cite: 1]
                    df2 = df2.drop_duplicates(subset=['Cycle', 'Prog Time'], keep='first') #[cite: 1]

                    df_voltage = df[['Cycle', 'Voltage']].copy() #[cite: 1]
                    final_table = pd.merge(final_table, df_voltage, on='Cycle', how='left') #[cite: 1]

                    # --- Getting Highest AhStep --- 
                    df3 = raw_df.copy() #[cite: 1]
                    df3['Cycle'] = pd.to_numeric(df3['Cycle'], errors='coerce') #[cite: 1]
                    df3['AhStep'] = pd.to_numeric(df3['AhStep'], errors='coerce') #[cite: 1]
                    df3_max = df3.groupby('Cycle', as_index=False)['AhStep'].max().rename(columns={'AhStep': '5 Hours Ah'}) #[cite: 1]

                    final_table = pd.merge(final_table, df3_max, on='Cycle', how='left') #[cite: 1]

                    # --- Getting Peak Voltage --- 
                    df4 = raw_df.copy() #[cite: 1]
                    df4['Cycle'] = pd.to_numeric(df4['Cycle'], errors='coerce') #[cite: 1]
                    df4['Voltage'] = pd.to_numeric(df4['Voltage'], errors='coerce') #[cite: 1]
                    df4_max = df4.groupby('Cycle', as_index=False)['Voltage'].max().rename(columns={'Voltage': 'Peak Voltage'}) #[cite: 1]

                    final_table = pd.merge(final_table, df4_max, on='Cycle', how='left') #[cite: 1]
                    
                    # --- Write to sheets --- 
                    df.to_excel(writer, sheet_name='Sheet 1', index=False) #[cite: 1]
                    df2.to_excel(writer, sheet_name='Sheet 2', index=False) #[cite: 1]
                    final_table.to_excel(writer, sheet_name='Sheet 3', index=False) #[cite: 1]

                    worksheet1 = writer.sheets['Sheet 1'] #[cite: 1]
                    worksheet2 = writer.sheets['Sheet 2'] #[cite: 1]
                    worksheet3 = writer.sheets['Sheet 3'] #[cite: 1]

                    # --- DoD 50% Charting Layout --- 
                    # Chart 1
                    chart1 = workbook.add_chart({'type': 'scatter', 'subtype': 'smooth_with_markers'}) #[cite: 1]
                    chart1.add_series({ #[cite: 1]
                        'categories': ['Sheet 3', 1, 0, len(final_table), 0], #[cite: 1]
                        'values':     ['Sheet 3', 1, 2, len(final_table), 2], #[cite: 1]
                        'line' :      {'width': 0.1, 'color': 'blue'} #[cite: 1]
                    }) #[cite: 1]

                    chart1.set_title({'name': f'{battery_code} - 50% DoD Endurance'}) #[cite: 1]
                    chart1.set_x_axis({'name': 'Cycle', 'label_position': 'low'}) #[cite: 1]
                    chart1.set_y_axis({'name': 'Voltage', 'min': 9.8, 'max': 12.2}) #[cite: 1]
                    chart1.set_size({'width': 750, 'height': 300}) #[cite: 1]
                    
                    worksheet3.insert_chart('G2', chart1) #[cite: 1]

                    # Chart 2
                    chart2 = workbook.add_chart({'type': 'scatter', 'subtype': 'smooth_with_markers'}) #[cite: 1]
                    chart2.add_series({ #[cite: 1]
                        'name':       'Voltage',  #[cite: 1]
                        'categories': ['Sheet 3', 1, 0, len(final_table), 0],  #[cite: 1]
                        'values':     ['Sheet 3', 1, 2, len(final_table), 2], #[cite: 1]
                        'line' :      {'width': 0.1, 'color': 'blue'} #[cite: 1]
                    }) #[cite: 1]

                    chart2.add_series({ #[cite: 1]
                        'name':       'Charged Time', #[cite: 1]
                        'categories': ['Sheet 3', 1, 0, len(final_table), 0],  #[cite: 1]
                        'values':     ['Sheet 3', 1, 1, len(final_table), 1], #[cite: 1]
                        'line' :      {'width': 0.1, 'color': 'red'}, #[cite: 1]
                        'y2_axis':    True   #[cite: 1]
                    }) #[cite: 1]

                    chart2.set_title({'name': f'{battery_code} - Discharged EOV vs Charged Time'}) #[cite: 1]
                    chart2.set_x_axis({'name': 'Cycle', 'label_position': 'low'}) #[cite: 1]
                    chart2.set_y_axis({'name': 'Voltage', 'min': 9.8, 'max': 12.2}) #[cite: 1]
                    chart2.set_y2_axis({'name': 'Charged Time (hrs)', 'min': 2, 'max': 6}) #[cite: 1]
                    chart2.set_size({'width': 750, 'height': 300}) #[cite: 1]
                    
                    worksheet3.insert_chart('G18', chart2) #[cite: 1]

                    # Chart 3
                    chart3 = workbook.add_chart({'type': 'scatter', 'subtype': 'smooth_with_markers'}) #[cite: 1]
                    
                    chart3.add_series({ #[cite: 1]
                        'name':       'Time',  #[cite: 1]
                        'categories': ['Sheet 3', 1, 0, len(final_table), 0],  #[cite: 1]
                        'values':     ['Sheet 3', 1, 1, len(final_table), 1],  #[cite: 1]
                        'line' :      {'width': 0.1, 'color': '#4472C4'},      #[cite: 1]
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#4472C4'}, 'fill': {'color': '#4472C4'}} #[cite: 1]
                    }) #[cite: 1]

                    chart3.add_series({ #[cite: 1]
                        'name':       '5hrs Ah', #[cite: 1]
                        'categories': ['Sheet 3', 1, 0, len(final_table), 0],  #[cite: 1]
                        'values':     ['Sheet 3', 1, 3, len(final_table), 3],  #[cite: 1]
                        'line' :      {'width': 0.1, 'color': '#ED7D31'},      #[cite: 1]
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#ED7D31'}, 'fill': {'color': '#ED7D31'}}, #[cite: 1]
                        'y2_axis':    True   #[cite: 1]
                    }) #[cite: 1]

                    chart3.set_title({'name': 'Chg Time Vs 5hrs Ah in'}) #[cite: 1]
                    chart3.set_x_axis({'name': 'Cycle', 'label_position': 'low'}) #[cite: 1]
                    chart3.set_y_axis({'name': 'Chg Time in hrs', 'min': 3.50, 'max': 5.50}) #[cite: 1]
                    chart3.set_y2_axis({'name': '5 hrs chg Ah', 'min': 31.00, 'max': 32.60}) #[cite: 1]
                    chart3.set_size({'width': 750, 'height': 300}) #[cite: 1]
                    
                    worksheet3.insert_chart('S2', chart3) #[cite: 1]

                    # Chart 4
                    chart4 = workbook.add_chart({'type': 'scatter', 'subtype': 'smooth_with_markers'}) #[cite: 1]
                    
                    chart4.add_series({ #[cite: 1]
                        'name':       'Time',  #[cite: 1]
                        'categories': ['Sheet 3', 1, 0, len(final_table), 0],  #[cite: 1]
                        'values':     ['Sheet 3', 1, 1, len(final_table), 1],  #[cite: 1]
                        'line' :      {'width': 0.1, 'color': '#4472C4'},      #[cite: 1]
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#4472C4'}, 'fill': {'color': '#4472C4'}} #[cite: 1]
                    }) #[cite: 1]

                    chart4.add_series({ #[cite: 1]
                        'name':       'Peak Voltage', #[cite: 1]
                        'categories': ['Sheet 3', 1, 0, len(final_table), 0],  #[cite: 1]
                        'values':     ['Sheet 3', 1, 4, len(final_table), 4],  #[cite: 1]
                        'line' :      {'width': 0.1, 'color': '#A5A5A5'},      #[cite: 1]
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#A5A5A5'}, 'fill': {'color': '#A5A5A5'}}, #[cite: 1]
                        'y2_axis':    True   #[cite: 1]
                    }) #[cite: 1]

                    chart4.set_title({'name': 'Chg Time Vs Peak Voltage'}) #[cite: 1]
                    chart4.set_x_axis({'name': 'Cycle', 'label_position': 'low'}) #[cite: 1]
                    chart4.set_y_axis({'name': 'Chg Time in hrs', 'min': 3.50, 'max': 5.50}) #[cite: 1]
                    chart4.set_y2_axis({'name': 'Peak Voltage V', 'min': 15.00, 'max': 18.00}) #[cite: 1]
                    chart4.set_size({'width': 750, 'height': 300}) #[cite: 1]
                    
                    worksheet3.insert_chart('S18', chart4) #[cite: 1]

                # Test D: Other types
                else:
                    st.warning(f"Bypassing custom calculations: Visual sheets layout for '{test_type}' hasn't been defined yet.")
                    st.stop()

            # --- Streamlit Download Button ---
            st.success(f"✅ Data processed successfully for {test_type} test!")
            
            st.download_button(
                label=f"📥 Download {output_filename}",
                data=output_buffer.getvalue(),
                file_name=output_filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"An error occurred while processing the file: {e}")
